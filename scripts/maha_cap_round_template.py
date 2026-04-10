"""
Build a long-format template: last N admission years × CAP rounds 1..R.

Example rows use Vishwakarma Institute of Information Technology (VIIT), Pune —
one sample branch + category. Closing merit / percentile are LEFT BLANK:
fill from official CET Cell CAP cutoff PDFs for each year.

Also writes a wide-format helper (one row per year, columns CAP1..CAPR) with
empty cells for pasting round closings for the same branch/category.
"""

from __future__ import annotations

import csv
from pathlib import Path

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None  # type: ignore[misc, assignment]

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"

# Five most recent completed / current-style admission cycles (adjust as needed).
ACADEMIC_YEARS = [
    "2021-22",
    "2022-23",
    "2023-24",
    "2024-25",
    "2025-26",
]
CAP_ROUNDS = [1, 2, 3, 4, 5]

LONG_HEADERS = [
    "Academic_Year",
    "CAP_Round",
    "College_Name",
    "Institute_Identifier_Notes",
    "Program_Branch",
    "Seat_Category_Code",
    "Closing_Merit_No",
    "Closing_Percentile",
    "Official_Source_URL",
    "Notes",
]

EXAMPLE_COLLEGE = (
    "Vishwakarma Institute of Information Technology (VIIT), Pune — "
    "match spelling & choice code to each year’s seat matrix / CAP PDF"
)
EXAMPLE_BRANCH = "B.Tech Computer Science and Engineering (example — duplicate template per branch)"
EXAMPLE_CATEGORY = "GOPEN (example — duplicate per category from PDF)"


def long_rows() -> list[list[str]]:
    rows: list[list[str]] = []
    for year in ACADEMIC_YEARS:
        for rnd in CAP_ROUNDS:
            rows.append(
                [
                    year,
                    str(rnd),
                    EXAMPLE_COLLEGE,
                    "Institute / course code: fill from that year’s seat matrix PDF",
                    EXAMPLE_BRANCH,
                    EXAMPLE_CATEGORY,
                    "",
                    "",
                    "",
                    "Paste cutoff from official CAP cutoff PDF; add document link in Official_Source_URL",
                ]
            )
    return rows


def write_long_csv(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(LONG_HEADERS)
        w.writerows(rows)


def write_wide_csv(path: Path, long_data: list[list[str]]) -> None:
    """One row per academic year; CAP1..CAP5 columns for closing merit (fill)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    # Map year -> {round: merit}
    by_year: dict[str, dict[int, str]] = {}
    for r in long_data:
        y, rnd_s, *_ = r[0], r[1], r[2]
        by_year.setdefault(y, {})[int(rnd_s)] = r[6]  # Closing_Merit_No

    headers = [
        "Academic_Year",
        "College_Name",
        "Program_Branch",
        "Seat_Category_Code",
        "CAP1_Closing_Merit",
        "CAP2_Closing_Merit",
        "CAP3_Closing_Merit",
        "CAP4_Closing_Merit",
        "CAP5_Closing_Merit",
        "Official_Source_URLs",
    ]
    rows_out: list[list[str]] = []
    for year in ACADEMIC_YEARS:
        rd = by_year.get(year, {})
        rows_out.append(
            [
                year,
                EXAMPLE_COLLEGE,
                EXAMPLE_BRANCH,
                EXAMPLE_CATEGORY,
                rd.get(1, ""),
                rd.get(2, ""),
                rd.get(3, ""),
                rd.get(4, ""),
                rd.get(5, ""),
                "",
            ]
        )
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows_out)


def write_xlsx(long_rows_data: list[list[str]], dest: Path) -> None:
    if Workbook is None:
        return
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    rm = wb.create_sheet("00_READ_ME", 0)
    for line in [
        "MAHARASHTRA-STYLE CAP ROUND HISTORY TEMPLATE",
        "",
        f"Sheets: long format ({len(long_rows_data)} rows = {len(ACADEMIC_YEARS)} years × {len(CAP_ROUNDS)} rounds),",
        "and wide format (one row per year, CAP1..CAP5 columns).",
        "",
        "Replace example VIIT / branch / category with your college and copy values",
        "from official CET Cell cutoff PDFs per year. Do not invent merit numbers.",
        "",
        "See reference/MAHARASHTRA_CAP_cutoff_sources.txt for where PDFs usually appear.",
    ]:
        rm.append([line])

    ws_long = wb.create_sheet("01_long_year_round")
    ws_long.append(LONG_HEADERS)
    for r in long_rows_data:
        ws_long.append(r)

    # Wide sheet
    by_year: dict[str, dict[int, str]] = {}
    for r in long_rows_data:
        y, rnd_s = r[0], r[1]
        by_year.setdefault(y, {})[int(rnd_s)] = r[6]

    ws_wide = wb.create_sheet("02_wide_CAP_columns")
    ws_wide.append(
        [
            "Academic_Year",
            "College_Name",
            "Program_Branch",
            "Seat_Category_Code",
            "CAP1_Closing_Merit",
            "CAP2_Closing_Merit",
            "CAP3_Closing_Merit",
            "CAP4_Closing_Merit",
            "CAP5_Closing_Merit",
            "Official_Source_URLs",
        ]
    )
    for year in ACADEMIC_YEARS:
        rd = by_year.get(year, {})
        ws_wide.append(
            [
                year,
                EXAMPLE_COLLEGE,
                EXAMPLE_BRANCH,
                EXAMPLE_CATEGORY,
                rd.get(1, ""),
                rd.get(2, ""),
                rd.get(3, ""),
                rd.get(4, ""),
                rd.get(5, ""),
                "",
            ]
        )

    wb.save(dest)


def main() -> None:
    rows = long_rows()
    long_path = DATA_DIR / "maha_cap_round_history_long_TEMPLATE.csv"
    wide_path = DATA_DIR / "maha_cap_round_history_wide_TEMPLATE.csv"
    xlsx_path = DATA_DIR / "maha_cap_round_history_TEMPLATE.xlsx"

    write_long_csv(long_path, rows)
    write_wide_csv(wide_path, rows)
    if Workbook is not None:
        write_xlsx(rows, xlsx_path)
        print(f"Wrote {xlsx_path}")
    else:
        print("openpyxl not installed — skip XLSX. pip install openpyxl")

    print(f"Wrote {long_path} ({len(rows)} rows)")
    print(f"Wrote {wide_path}")


if __name__ == "__main__":
    main()
