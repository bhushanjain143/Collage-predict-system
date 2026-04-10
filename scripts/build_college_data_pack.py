"""
Build engineering college data pack: CSV + XLSX (+ optional ZIP).
Institute names/locations for IIT/NIT/IIIT tiers are factual labels only;
cutoff, fees, placements, and scholarships are NOT scraped — placeholders
point users to official JoSAA / institute sources.
"""

from __future__ import annotations

import argparse
import csv
import zipfile
from pathlib import Path

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None  # type: ignore[misc, assignment]

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUT_XLSX = DATA_DIR / "engineering_colleges_SAMPLE_JoSAA_tier.xlsx"
OUT_ZIP = ROOT / "dist" / "engineering_colleges_data_pack.zip"

HEADERS = [
    "Institute_Type",
    "College Name",
    "City",
    "Branch",
    "Cutoff",
    "Fees",
    "Placement %",
    "Avg Package",
    "Highest Package",
    "Scholarship",
]

# Placeholders — do not treat as verified admission or placement statistics.
PH_BRANCH = (
    "All UG programmes (duplicate row per branch from official brochure / seat matrix)"
)
PH_CUTOFF = (
    "See JoSAA/CSAB closing ranks (program, category, round, year) — https://josaa.nic.in/"
)
PH_FEES = "See official institute fee notification (tentative until confirmed with college)"
PH_PLACE = "See institute placement report / disclosure (definition varies)"
PH_AVG = PH_PLACE
PH_HIGH = PH_PLACE
PH_SCH = "See institute scholarships + state/central schemes (eligibility varies)"

IITS: list[tuple[str, str]] = [
    ("Indian Institute of Technology Kharagpur", "Kharagpur"),
    ("Indian Institute of Technology Bombay", "Mumbai"),
    ("Indian Institute of Technology Madras", "Chennai"),
    ("Indian Institute of Technology Kanpur", "Kanpur"),
    ("Indian Institute of Technology Delhi", "New Delhi"),
    ("Indian Institute of Technology Guwahati", "Guwahati"),
    ("Indian Institute of Technology Roorkee", "Roorkee"),
    ("Indian Institute of Technology Ropar", "Rupnagar"),
    ("Indian Institute of Technology Bhubaneswar", "Bhubaneswar"),
    ("Indian Institute of Technology Gandhinagar", "Gandhinagar"),
    ("Indian Institute of Technology Hyderabad", "Hyderabad"),
    ("Indian Institute of Technology Jodhpur", "Jodhpur"),
    ("Indian Institute of Technology Patna", "Patna"),
    ("Indian Institute of Technology Indore", "Indore"),
    ("Indian Institute of Technology Mandi", "Mandi"),
    ("Indian Institute of Technology (BHU) Varanasi", "Varanasi"),
    ("Indian Institute of Technology Palakkad", "Palakkad"),
    ("Indian Institute of Technology Tirupati", "Tirupati"),
    ("Indian Institute of Technology Bhilai", "Bhilai"),
    ("Indian Institute of Technology Goa", "Ponda"),
    ("Indian Institute of Technology Dharwad", "Dharwad"),
    ("Indian Institute of Technology Jammu", "Jammu"),
    ("Indian Institute of Technology (ISM) Dhanbad", "Dhanbad"),
]

NITS: list[tuple[str, str]] = [
    ("National Institute of Technology Agartala", "Agartala"),
    ("National Institute of Technology Andhra Pradesh", "Tadepalligudem"),
    ("National Institute of Technology Arunachal Pradesh", "Yupia"),
    ("National Institute of Technology Calicut", "Kozhikode"),
    ("National Institute of Technology Delhi", "New Delhi"),
    ("National Institute of Technology Durgapur", "Durgapur"),
    ("National Institute of Technology Goa", "Ponda"),
    ("National Institute of Technology Hamirpur", "Hamirpur"),
    ("National Institute of Technology Karnataka, Surathkal", "Mangaluru"),
    ("National Institute of Technology Kurukshetra", "Kurukshetra"),
    ("National Institute of Technology Manipur", "Imphal"),
    ("National Institute of Technology Meghalaya", "Shillong"),
    ("National Institute of Technology Mizoram", "Aizawl"),
    ("National Institute of Technology Nagaland", "Chümoukedima"),
    ("National Institute of Technology Patna", "Patna"),
    ("National Institute of Technology Puducherry", "Karaikal"),
    ("National Institute of Technology Raipur", "Raipur"),
    ("National Institute of Technology Rourkela", "Rourkela"),
    ("National Institute of Technology Sikkim", "Ravangla"),
    ("National Institute of Technology Silchar", "Silchar"),
    ("National Institute of Technology Srinagar", "Srinagar"),
    ("National Institute of Technology, Surat", "Surat"),
    ("National Institute of Technology Tiruchirappalli", "Tiruchirappalli"),
    ("National Institute of Technology Uttarakhand", "Srinagar (Garhwal)"),
    ("National Institute of Technology Warangal", "Warangal"),
    ("Dr. B. R. Ambedkar National Institute of Technology Jalandhar", "Jalandhar"),
    ("Malaviya National Institute of Technology Jaipur", "Jaipur"),
    ("Maulana Azad National Institute of Technology Bhopal", "Bhopal"),
    ("Motilal Nehru National Institute of Technology Allahabad", "Prayagraj"),
    ("National Institute of Technology Jamshedpur", "Jamshedpur"),
    ("Visvesvaraya National Institute of Technology Nagpur", "Nagpur"),
]

# Statutory IIITs — aligned with MHRD "list_iiits.pdf" (25 orgs: 5 MOE + 20 PPP).
# Some JoSAA tables quote 26 IIIT-class institutes; verify your year on josaa.nic.in.
IIITS: list[tuple[str, str]] = [
    (
        "Atal Bihari Vajpayee Indian Institute of Information Technology and Management, Gwalior",
        "Gwalior",
    ),
    ("Indian Institute of Information Technology, Allahabad", "Prayagraj"),
    (
        "Indian Institute of Information Technology, Design and Manufacturing, Jabalpur",
        "Jabalpur",
    ),
    (
        "Indian Institute of Information Technology, Design and Manufacturing, Kancheepuram",
        "Kancheepuram",
    ),
    ("Indian Institute of Information Technology, Sri City", "Sri City"),
    ("Indian Institute of Information Technology, Guwahati", "Guwahati"),
    ("Indian Institute of Information Technology, Vadodara", "Vadodara"),
    ("Indian Institute of Information Technology, Kota", "Kota"),
    ("Indian Institute of Information Technology, Tiruchirappalli", "Tiruchirappalli"),
    ("Indian Institute of Information Technology, Dharwad", "Dharwad"),
    ("Indian Institute of Information Technology, Una", "Una"),
    ("Indian Institute of Information Technology, Sonepat", "Sonipat"),
    ("Indian Institute of Information Technology, Kalyani", "Kalyani"),
    ("Indian Institute of Information Technology, Lucknow", "Lucknow"),
    (
        "Indian Institute of Information Technology, Design and Manufacturing, Kurnool",
        "Kurnool",
    ),
    ("Indian Institute of Information Technology, Kottayam", "Kottayam"),
    ("Indian Institute of Information Technology, Manipur", "Imphal"),
    ("Indian Institute of Information Technology, Nagpur", "Nagpur"),
    ("Indian Institute of Information Technology, Pune", "Pune"),
    ("Indian Institute of Information Technology, Ranchi", "Ranchi"),
    ("Indian Institute of Information Technology, Surat", "Surat"),
    ("Indian Institute of Information Technology, Bhopal", "Bhopal"),
    ("Indian Institute of Information Technology, Bhagalpur", "Bhagalpur"),
    ("Indian Institute of Information Technology, Agartala", "Agartala"),
    ("Indian Institute of Information Technology, Raichur", "Raichur"),
]


def rows_for_tier(pairs: list[tuple[str, str]], institute_type: str) -> list[list[str]]:
    out: list[list[str]] = []
    for name, city in pairs:
        out.append(
            [
                institute_type,
                name,
                city,
                PH_BRANCH,
                PH_CUTOFF,
                PH_FEES,
                PH_PLACE,
                PH_AVG,
                PH_HIGH,
                PH_SCH,
            ]
        )
    return out


def write_csv(path: Path, rows: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(rows)


def write_xlsx(
    path: Path,
    combined: list[list[str]],
    iit_rows: list[list[str]],
    nit_rows: list[list[str]],
    iiit_rows: list[list[str]],
) -> None:
    if Workbook is None:
        raise RuntimeError("openpyxl is required for XLSX export: pip install openpyxl")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    readme = wb.create_sheet("00_READ_ME", 0)
    readme_notes = [
        "ENGINEERING COLLEGE SAMPLE PACK — READ FIRST",
        "",
        "Sheet '01_Combined_IIT_NIT_IIIT' = all 79 institutes in one table (23 IIT + 31 NIT + 25 IIIT).",
        "Sheets '02_IIT_only', '03_NIT_only', '04_IIIT_only' = same columns, split by type so you can see each group clearly.",
        "",
        "Institute_Type column: IIT | NIT | IIIT (JoSAA counselling tier).",
        "This file does NOT list 6,611+ private or 2,265+ government colleges — use AICTE approved-institute export; see reference/AICTE_bulk_college_list_howto.txt in the ZIP.",
        "",
        "IIIT count: MHRD list_iiits.pdf lists 25 statutory IIIT organisations (5 MOE + 20 PPP).",
        "Some JoSAA materials quote 26 IIIT-class seats/institutes for a year — verify on josaa.nic.in seat matrix for your admission year.",
        "",
        "Cutoff / fees / placement / scholarship cells are placeholders — fill from official JoSAA PDFs + institute sites.",
    ]
    for line in readme_notes:
        readme.append([line])

    def fill_sheet(title: str, rows: list[list[str]]) -> None:
        ws = wb.create_sheet(title)
        ws.append(HEADERS)
        for r in rows:
            ws.append(r)

    fill_sheet("01_Combined_IIT_NIT_IIIT", combined)
    fill_sheet("02_IIT_only", iit_rows)
    fill_sheet("03_NIT_only", nit_rows)
    fill_sheet("04_IIIT_only", iiit_rows)
    wb.save(path)


def write_zip(
    csv_path: Path,
    xlsx_path: Path | None,
    template_csv: Path,
    guide_path: Path,
    reference_dir: Path,
    dest: Path,
) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(csv_path, arcname=f"data/{csv_path.name}")
        if xlsx_path is not None and xlsx_path.exists():
            zf.write(xlsx_path, arcname=f"data/{xlsx_path.name}")
        if template_csv.exists():
            zf.write(template_csv, arcname=f"data/{template_csv.name}")
        if guide_path.exists():
            zf.write(guide_path, arcname=guide_path.name)
        if reference_dir.is_dir():
            for p in sorted(reference_dir.glob("*.txt")):
                zf.write(p, arcname=f"reference/{p.name}")
        for p in sorted(DATA_DIR.glob("maha_cap*")):
            if p.is_file():
                zf.write(p, arcname=f"data/{p.name}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-xlsx", action="store_true", help="Skip Excel if openpyxl missing")
    parser.add_argument("--no-zip", action="store_true")
    args = parser.parse_args()

    iit_rows = rows_for_tier(IITS, "IIT")
    nit_rows = rows_for_tier(NITS, "NIT")
    iiit_rows = rows_for_tier(IIITS, "IIIT")
    rows: list[list[str]] = []
    rows.extend(iit_rows)
    rows.extend(nit_rows)
    rows.extend(iiit_rows)

    csv_path = DATA_DIR / "engineering_colleges_SAMPLE_JoSAA_tier.csv"
    write_csv(csv_path, rows)

    guide = ROOT / "DATA_COLLECTION_GUIDE.txt"
    template_csv = DATA_DIR / "engineering_colleges_TEMPLATE.csv"
    reference_dir = ROOT / "reference"

    xlsx_written: Path | None = None
    if not args.no_xlsx and Workbook is not None:
        write_xlsx(OUT_XLSX, rows, iit_rows, nit_rows, iiit_rows)
        xlsx_written = OUT_XLSX
    elif not args.no_xlsx:
        print("openpyxl not installed — skipping XLSX. Run: pip install openpyxl")

    if not args.no_zip:
        write_zip(csv_path, xlsx_written, template_csv, guide, reference_dir, OUT_ZIP)
        print(f"Wrote {csv_path} ({len(rows)} data rows)")
        if OUT_XLSX.exists():
            print(f"Wrote {OUT_XLSX}")
        print(f"Wrote {OUT_ZIP}")


if __name__ == "__main__":
    main()
